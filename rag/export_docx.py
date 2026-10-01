"""Exportação da peça gerada para Word (.docx)."""
from __future__ import annotations

import io
import re
from typing import Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt

_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\[(?:PREENCHER|VERIFICAR)[^\]]*\])")


def _runs(paragrafo, texto: str) -> None:
    for parte in _INLINE_RE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            paragrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith(("[PREENCHER", "[VERIFICAR")):
            run = paragrafo.add_run(parte)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            paragrafo.add_run(parte)


def _corpo(doc, texto: str) -> None:
    for linha in texto.split("\n"):
        linha = linha.strip()
        if not linha:
            continue
        if re.match(r"^#{1,4}\s", linha):
            p = doc.add_paragraph()
            p.add_run(linha.lstrip("#").strip()).bold = True
        elif re.match(r"^[-*•]\s", linha):
            _runs(doc.add_paragraph(style="List Bullet"), linha[2:].strip())
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(p, linha)


def exportar_docx(
    titulo: str,
    subtitulo: str,
    numeracao: str,
    secoes: List[Dict[str, str]],
    fontes: Optional[List[str]] = None,
) -> bytes:
    """`secoes`: lista de dicts com 'rotulo', 'titulo', 'texto' e opcionalmente 'parte'."""
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)

    h = doc.add_heading(titulo.upper(), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitulo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(subtitulo).italic = True

    parte_atual = None
    for s in secoes:
        if s.get("parte") and s["parte"] != parte_atual:
            parte_atual = s["parte"]
            ph = doc.add_heading(parte_atual, level=1)
            ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if numeracao in ("artigo", "clausula"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.add_run(s["rotulo"]).bold = True
            p.add_run("\n")
            p.add_run(s["titulo"]).bold = True
        else:
            doc.add_heading(f"{s['rotulo']} {s['titulo']}", level=2)
        _corpo(doc, s.get("texto", ""))

    if fontes:
        doc.add_page_break()
        doc.add_heading("Anexo de rastreabilidade — fontes consultadas (uso interno, remover antes de publicar)", level=2)
        for f in fontes:
            doc.add_paragraph(f, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
